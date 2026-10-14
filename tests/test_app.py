import docx
import pytest
import yaml
from streamlit.testing.v1 import AppTest

from conftest import ROOT
from cornerstone.cli import DB_FILE, INDEX_FILE
from cornerstone.embeddings import Embedder
from cornerstone.index.hnsw_index import HNSWIndex
from cornerstone.index.metadata_store import MetadataStore
from cornerstone.pipeline import ingest_file

APP = str(ROOT / "app" / "streamlit_app.py")


@pytest.fixture(scope="module")
def prepared(cfg, tmp_path_factory):
    """A tiny real index in a temp folder + a config file pointing at it."""
    tmp = tmp_path_factory.mktemp("app")
    c = {**cfg, "paths": {"raw_dir": str(tmp / "raw"), "index_dir": str(tmp / "index")}}
    cfg_file = tmp / "config.yaml"
    cfg_file.write_text(yaml.safe_dump(c), encoding="utf-8")
    d = docx.Document()
    d.add_heading("Translation look-aside buffer", level=1)
    d.add_paragraph("A TLB is a small fast associative cache of page table entries used by the MMU. " * 5)
    f = tmp / "memory.docx"
    d.save(f)
    store = MetadataStore(tmp / "index" / DB_FILE)
    index = HNSWIndex.from_config(c)
    r = ingest_file(f, c, store, index, Embedder(c), tmp / "index" / INDEX_FILE)
    assert r.status == "indexed"
    store.close()
    return cfg_file


def run_app(cfg_file, monkeypatch):
    monkeypatch.setenv("CORNERSTONE_CONFIG", str(cfg_file))
    at = AppTest.from_file(APP, default_timeout=120)
    return at.run()


def test_page_renders_with_stats_and_no_exception(prepared, monkeypatch):
    at = run_app(prepared, monkeypatch)
    assert not at.exception
    assert at.title[0].value == "Cornerstone"
    labels = {m.label: m.value for m in at.sidebar.metric}
    assert labels["Documents"] == "1" and int(labels["Chunks (searchable)"]) >= 1


def test_search_shows_result_with_source(prepared, monkeypatch):
    at = run_app(prepared, monkeypatch)
    at.text_input[0].set_value("What is a TLB?")
    at.button[0].click()          # the form's submit button
    at.run()
    assert not at.exception
    text = " ".join(m.value for m in at.markdown)
    assert "memory.docx" in text and "score" in text and "paragraphs" in text


def test_blank_query_shows_message_not_crash(prepared, monkeypatch):
    at = run_app(prepared, monkeypatch)
    at.text_input[0].set_value("   ")
    at.button[0].click()
    at.run()
    assert not at.exception
    assert any("type a question" in e.value for e in at.error)


def test_upload_ingest_report_and_duplicate(cfg, tmp_path, monkeypatch):
    """Upload a generated DOCX through the page, check the report, then upload it again."""
    import io
    c = {**cfg, "paths": {"raw_dir": str(tmp_path / "raw"), "index_dir": str(tmp_path / "index")}}
    cfg_file = tmp_path / "config.yaml"
    cfg_file.write_text(yaml.safe_dump(c), encoding="utf-8")
    d = docx.Document()
    d.add_heading("Scheduling", level=1)
    d.add_paragraph("Round robin gives each process a fixed time quantum in turn. " * 6)
    buf = io.BytesIO()
    d.save(buf)
    payload = ("sched.docx", buf.getvalue(),
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document")

    monkeypatch.setenv("CORNERSTONE_CONFIG", str(cfg_file))
    at = AppTest.from_file(APP, default_timeout=180).run()
    assert not at.exception
    at.file_uploader[0].set_value([payload])
    at.run()
    ingest_btn = next(b for b in at.button if b.label == "Ingest")
    ingest_btn.click()
    at.run()
    assert not at.exception
    rep = at.dataframe[0].value.iloc[0]
    assert rep["file"] == "sched.docx" and rep["status"] == "indexed" and rep["chunks"] >= 1
    assert (tmp_path / "raw" / "sched.docx").is_file()

    at.run()                                                   # fresh run: index reloaded from disk
    assert {m.label: m.value for m in at.sidebar.metric}["Documents"] == "1"
    at.file_uploader[0].set_value([payload])
    at.run()
    next(b for b in at.button if b.label == "Ingest").click()
    at.run()
    assert at.dataframe[0].value.iloc[0]["status"] == "skipped_duplicate"