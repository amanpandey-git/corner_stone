import docx
import pytest

from cornerstone import EmptyQuery
from cornerstone.embeddings import Embedder
from cornerstone.index.hnsw_index import HNSWIndex
from cornerstone.index.metadata_store import MetadataStore
from cornerstone.pipeline import ingest_file
from cornerstone.search import search

from conftest import SAMPLES


@pytest.fixture(scope="module")
def embedder(cfg):
    return Embedder(cfg)


def new_stack(cfg, tmp_path):
    store = MetadataStore(tmp_path / "meta.sqlite")
    return store, HNSWIndex.from_config(cfg), tmp_path / "hnsw.faiss"


def make_docx(path, topic, n=10):
    d = docx.Document()
    d.add_heading(f"{topic.title()}", level=1)
    for i in range(n):
        d.add_paragraph(f"Paragraph {i}: {topic} is discussed here in detail with examples number {i}. " * 3)
    d.save(path)
    return path


@pytest.fixture(scope="module")
def samples_stack(cfg, embedder, tmp_path_factory):
    tmp = tmp_path_factory.mktemp("samples_idx")
    store, index, ipath = new_stack(cfg, tmp)
    for name in ("Unit 5_Part 2_System Protection.pdf", "Unit 4_Storage management_ Disk Schduling.pdf"):
        r = ingest_file(SAMPLES / name, cfg, store, index, embedder, ipath)
        assert r.status == "indexed"
    return store, index


def test_known_questions_return_expected_file(cfg, embedder, samples_stack):
    store, index = samples_stack
    cases = [
        ("What is the principle of least privilege?", "Unit 5_Part 2_System Protection.pdf"),
        ("Explain the SCAN elevator disk scheduling algorithm", "Unit 4_Storage management_ Disk Schduling.pdf"),
    ]
    for q, expected in cases:
        res = search(q, 5, cfg, store, index, embedder)
        assert res[0].chunk.file_name == expected, (q, [r.chunk.file_name for r in res])


def test_result_fields_and_ordering(cfg, embedder, samples_stack):
    store, index = samples_stack
    res = search("access matrix and domains of protection", 5, cfg, store, index, embedder)
    assert [r.rank for r in res] == [1, 2, 3, 4, 5]
    assert all(res[i].score >= res[i + 1].score for i in range(4))
    for r in res:
        c = r.chunk
        assert c.file_name.endswith(".pdf") and c.page_start is not None and c.page_end >= c.page_start
        assert c.chunk_id > 0 and c.token_count > 0 and c.text
        assert 0 < len(r.snippet) <= cfg["search"]["snippet_chars"] + 3
        assert "\n" not in r.snippet


def test_replaced_document_old_chunks_never_appear(cfg, embedder, tmp_path):
    store, index, ipath = new_stack(cfg, tmp_path)
    f = make_docx(tmp_path / "notes.docx", "round robin cpu scheduling quantum")
    r1 = ingest_file(f, cfg, store, index, embedder, ipath)
    old_ids = set(store.active_chunk_ids())
    assert search("round robin scheduling quantum", 5, cfg, store, index, embedder)
    make_docx(f, "disk head seek time and rotational latency")
    r2 = ingest_file(f, cfg, store, index, embedder, ipath)
    assert (r1.status, r2.status) == ("indexed", "replaced")
    for q in ("round robin scheduling quantum", "disk seek time", "anything at all"):
        res = search(q, 10, cfg, store, index, embedder)
        assert res and not ({r.chunk.chunk_id for r in res} & old_ids)
        assert all("round robin" not in r.chunk.text for r in res)
    store.close()


def test_deleted_document_never_appears_and_all_deleted_returns_empty(cfg, embedder, tmp_path):
    store, index, ipath = new_stack(cfg, tmp_path)
    a = ingest_file(make_docx(tmp_path / "a.docx", "semaphore mutex locking"), cfg, store, index, embedder, ipath)
    b = ingest_file(make_docx(tmp_path / "b.docx", "virtual memory page fault"), cfg, store, index, embedder, ipath)
    store.mark_replaced(a.doc_id)
    res = search("semaphore mutex locking", 5, cfg, store, index, embedder)
    assert res and all(r.chunk.doc_id == b.doc_id for r in res)   # a's vectors are closer but filtered
    store.mark_replaced(b.doc_id)
    assert search("semaphore mutex locking", 5, cfg, store, index, embedder) == []
    store.close()


def test_overfetch_finds_live_results_behind_many_deleted_ones(cfg, embedder, tmp_path):
    store, index, ipath = new_stack(cfg, tmp_path)
    old = ingest_file(make_docx(tmp_path / "old.docx", "deadlock detection and recovery", n=30),
                      cfg, store, index, embedder, ipath)
    live = ingest_file(make_docx(tmp_path / "live.docx", "deadlock prevention banker algorithm", n=6),
                       cfg, store, index, embedder, ipath)
    store.mark_replaced(old.doc_id)
    k = len(store.active_chunk_ids())                          # live doc is small: a few chunks
    assert old.chunks_created > 3 * k                          # first fetch (3k) is all deleted chunks
    res = search("deadlock detection and recovery", k, cfg, store, index, embedder)
    assert len(res) == k and all(r.chunk.doc_id == live.doc_id for r in res)
    store.close()


def test_empty_index_k_too_large_and_bad_inputs(cfg, embedder, tmp_path):
    store, index, ipath = new_stack(cfg, tmp_path)
    assert search("anything", 5, cfg, store, index, embedder) == []
    with pytest.raises(EmptyQuery):
        search("   ", 5, cfg, store, index, embedder)
    r = ingest_file(make_docx(tmp_path / "s.docx", "cache memory", n=1), cfg, store, index, embedder, ipath)
    res = search("cache", 50, cfg, store, index, embedder)
    assert len(res) == r.chunks_created >= 1
    with pytest.raises(ValueError):
        search("cache", 0, cfg, store, index, embedder)
    store.close()


def test_collapse_switch_hides_near_copies_across_files(cfg, embedder, tmp_path):
    store, index, ipath = new_stack(cfg, tmp_path)
    for name, extra in (("notes_a.docx", "Written in week one."), ("notes_b.docx", "Copied for revision.")):
        d = docx.Document()
        d.add_heading("Deadlock", level=1)
        d.add_paragraph("A deadlock happens when every process in a set waits for a resource held by another "
                        "process in the set, so none of them can ever continue. Four conditions must hold together: "
                        "mutual exclusion, hold and wait, no preemption and circular wait. The system can prevent "
                        "deadlock by denying one condition, avoid it with the Banker's algorithm, or detect it and "
                        "recover by terminating processes or preempting resources. " + extra)
        d.save(tmp_path / name)
        assert ingest_file(tmp_path / name, cfg, store, index, embedder, ipath).status == "indexed"
    on = search("what is a deadlock", 5, cfg, store, index, embedder)
    off = search("what is a deadlock", 5, {**cfg, "retrieval": {**cfg["retrieval"], "collapse": False}},
                 store, index, embedder)
    assert len(off) == 2 and len(on) == 1
    assert on[0].also_on and on[0].also_on[0][0] in {"notes_a.docx", "notes_b.docx"}
    assert on[0].rank == 1 and on[0].scores["vector"] == pytest.approx(on[0].score)
