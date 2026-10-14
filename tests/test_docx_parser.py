import docx
import pytest

from cornerstone import CorruptFile, EmptyFile, PasswordProtected, UnsupportedFileType
from cornerstone.ingest.docx_parser import extract_docx
from cornerstone.ingest.loader import load_document


def _make_docx(path, spec):
    """spec: list of ('h1'|'h2'|'p'|'table', text)."""
    d = docx.Document()
    for kind, text in spec:
        if kind == "h1":
            d.add_heading(text, level=1)
        elif kind == "h2":
            d.add_heading(text, level=2)
        elif kind == "table":
            t = d.add_table(rows=1, cols=2)
            t.rows[0].cells[0].text, t.rows[0].cells[1].text = text.split("|")
        else:
            d.add_paragraph(text)
    d.save(path)
    return path


def test_sections_have_paragraph_ranges_and_no_page_numbers(tmp_path, cfg):
    f = _make_docx(tmp_path / "a.docx", [
        ("p", "Preface text before any heading."),
        ("h1", "Scheduling"), ("p", "FCFS is simple."), ("p", "SJF minimises waiting."),
        ("h2", "Round robin"), ("p", "Time quantum matters."),
    ])
    pages = extract_docx(f, cfg)
    assert [p.page_no for p in pages] == [None, None, None]
    assert pages[1].headings[0].text == "Scheduling" and pages[1].headings[0].level == 1
    assert pages[2].headings[0].level == 2
    assert pages[1].para_start == 1 and pages[1].para_end == 3
    assert pages[0].para_start == 0 and pages[0].para_end == 0


def test_no_headings_gives_single_section_and_tables_not_dropped(tmp_path, cfg):
    f = _make_docx(tmp_path / "b.docx", [("p", "Intro line."), ("table", "Algo|Complexity"), ("p", "End.")])
    pages = extract_docx(f, cfg)
    assert len(pages) == 1
    assert "Algo | Complexity" in pages[0].text
    assert pages[0].headings == []


def test_empty_document_is_flagged_not_skipped(tmp_path, cfg):
    f = _make_docx(tmp_path / "c.docx", [])
    pages = extract_docx(f, cfg)
    assert len(pages) == 1 and pages[0].status == "empty" and pages[0].reason


def test_empty_corrupt_encrypted(tmp_path, cfg):
    e = tmp_path / "e.docx"
    e.write_bytes(b"")
    with pytest.raises(EmptyFile):
        extract_docx(e, cfg)
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip at all" * 10)
    with pytest.raises(CorruptFile):
        extract_docx(bad, cfg)
    enc = tmp_path / "enc.docx"
    enc.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100)
    with pytest.raises(PasswordProtected):
        extract_docx(enc, cfg)


def test_loader_dispatch_and_unsupported(tmp_path, small_pdf, cfg):
    assert len(load_document(small_pdf, cfg)) == 28
    d = _make_docx(tmp_path / "x.docx", [("h1", "T"), ("p", "body words here")])
    assert load_document(d, cfg)[0].text.startswith("T")
    txt = tmp_path / "n.txt"
    txt.write_text("hi")
    with pytest.raises(UnsupportedFileType):
        load_document(txt, cfg)
