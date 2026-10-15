import shutil

import docx
import pymupdf
import pytest

from cornerstone.embeddings import Embedder
from cornerstone.index.hnsw_index import HNSWIndex
from cornerstone.index.metadata_store import MetadataStore
from cornerstone.pipeline import ingest_file
from conftest import SAMPLES


class Crash(BaseException):
    """Stands in for a killed process: not an Exception, so cleanup code does not run."""


@pytest.fixture(scope="module")
def embedder(cfg):
    return Embedder(cfg)


@pytest.fixture
def stack(cfg, embedder, tmp_path):
    store = MetadataStore(tmp_path / "idx" / "meta.sqlite")
    index = HNSWIndex.from_config(cfg)
    yield store, index, tmp_path / "idx" / "hnsw.faiss", tmp_path
    store.close()


def make_docx(path, topic):
    d = docx.Document()
    d.add_heading(f"{topic.title()} basics", level=1)
    d.add_paragraph(f"{topic} is an important operating system concept. " * 12)
    d.save(path)
    return path


def run(path, cfg, stack, embedder, **kw):
    store, index, idx_path, _ = stack
    return ingest_file(path, cfg, store, index, embedder, idx_path, **kw)


def test_ingest_sample_pdf(cfg, stack, embedder, small_pdf):
    store, index, idx_path, _ = stack
    msgs = []
    r = run(small_pdf, cfg, stack, embedder, progress=msgs.append)
    assert r.status == "indexed" and r.doc_id
    assert r.pages_total == 28 and r.pages_with_text + r.pages_flagged == 28
    assert r.pages_flagged == len(r.flagged_pages) and all("OCR" in why for _, why in r.flagged_pages)
    assert r.chunks_created > 10 and r.seconds > 0
    assert len(index) == r.chunks_created and idx_path.is_file()
    assert store.stats()["chunks_active"] == r.chunks_created
    assert store.check_sync(set(index.all_ids().tolist())) == []
    assert any("embedded" in m for m in msgs)


def test_duplicate_is_skipped_even_under_another_name(cfg, stack, embedder, small_pdf):
    store, index, _, tmp = stack
    first = run(small_pdf, cfg, stack, embedder)
    n = len(index)
    again = run(small_pdf, cfg, stack, embedder)
    assert again.status == "skipped_duplicate" and again.doc_id == first.doc_id and again.chunks_created == 0
    copy = tmp / "renamed.pdf"
    shutil.copy(small_pdf, copy)
    assert run(copy, cfg, stack, embedder).status == "skipped_duplicate"
    assert len(index) == n and store.stats()["documents"] == 1


def test_changed_content_same_name_replaces_old_version(cfg, stack, embedder):
    store, index, idx_path, tmp = stack
    f = make_docx(tmp / "notes.docx", "deadlock")
    r1 = run(f, cfg, stack, embedder)
    old_ids = store.active_chunk_ids()
    make_docx(f, "paging")                                     # same name, new content
    r2 = run(f, cfg, stack, embedder)
    assert (r1.status, r2.status) == ("indexed", "replaced")
    assert store.get_document(r1.doc_id).status == "replaced"
    assert store.find_active_by_name("notes.docx").doc_id == r2.doc_id
    new_ids = store.active_chunk_ids()
    assert new_ids and not set(new_ids) & set(old_ids)
    assert store.get_chunks(old_ids) == {}
    assert len(index) == len(old_ids) + len(new_ids)           # old vectors stay in graph, filtered by SQLite
    assert store.check_sync(set(index.all_ids().tolist())) == []


@pytest.mark.parametrize("kind", ["unsupported", "empty", "corrupt", "password"])
def test_bad_file_gives_failed_report_and_writes_nothing(cfg, stack, embedder, kind):
    store, index, idx_path, tmp = stack
    if kind == "unsupported":
        f = tmp / "notes.txt"
        f.write_text("hello")
    elif kind == "empty":
        f = tmp / "empty.pdf"
        f.write_bytes(b"")
    elif kind == "corrupt":
        f = tmp / "bad.pdf"
        f.write_bytes(b"not a pdf at all" * 30)
    else:
        src = tmp / "plain.pdf"
        d = pymupdf.open()
        d.new_page().insert_text((72, 100), "secret")
        d.save(src)
        f = tmp / "locked.pdf"
        d.save(f, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="o", user_pw="u")
        d.close()
    r = run(f, cfg, stack, embedder)
    assert r.status == "failed" and r.message
    assert {"unsupported": "UnsupportedFileType", "empty": "EmptyFile", "corrupt": "CorruptFile",
            "password": "PasswordProtected"}[kind] in r.message
    st = store.stats()
    assert st["documents"] == st["documents_pending"] == st["chunks_active"] == st["chunks_pending"] == 0
    assert store.conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 0
    assert len(index) == 0 and not idx_path.exists()


def test_missing_file_fails_cleanly(cfg, stack, embedder, tmp_path):
    r = run(tmp_path / "ghost.pdf", cfg, stack, embedder)
    assert r.status == "failed" and "cannot read" in r.message


def test_crash_after_add_document_leaves_pending_detected_and_cleanable(cfg, stack, embedder, monkeypatch):
    store, index, idx_path, tmp = stack
    f = make_docx(tmp / "crash.docx", "semaphore")

    def die(*a, **k):
        raise Crash()
    monkeypatch.setattr(index, "save", die)                    # process "dies" after vectors were added
    with pytest.raises(Crash):
        run(f, cfg, stack, embedder)
    doc = store.list_documents("pending")
    assert len(doc) == 1 and doc[0].file_name == "crash.docx"
    assert store.find_by_hash(doc[0].sha256) is None           # not considered ingested
    assert store.stats()["chunks_active"] == 0                 # invisible to search
    problems = store.check_sync(set(index.all_ids().tolist()))
    assert any("interrupted" in p for p in problems)
    store.discard_pending(doc[0].doc_id)
    assert store.check_sync(set(index.all_ids().tolist())) == []
    monkeypatch.undo()
    assert run(f, cfg, stack, embedder).status == "indexed"    # re-ingest works after cleanup


def test_ordinary_failure_after_add_cleans_up_itself(cfg, stack, embedder, monkeypatch):
    store, index, _, tmp = stack
    f = make_docx(tmp / "disk.docx", "cache")
    monkeypatch.setattr(index, "save", lambda p: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(OSError):
        run(f, cfg, stack, embedder)
    assert store.list_documents("pending") == []
    assert store.stats()["documents_abandoned"] == 1


def test_document_with_only_flagged_pages_is_stored_not_hidden(cfg, stack, embedder, tmp_path):
    store, _, _, tmp = stack
    p = tmp / "scan.pdf"
    d = pymupdf.open()
    page = d.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 40, 40), False)
    pix.clear_with(200)
    page.insert_image(pymupdf.Rect(50, 50, 300, 300), pixmap=pix)
    d.save(p)
    d.close()
    r = run(p, cfg, stack, embedder)
    assert r.status == "indexed" and r.chunks_created == 0 and r.pages_flagged == 1
    assert "no indexable text" in r.message
    assert store.stats()["flagged_pages"][0]["status"] == "needs_ocr"


def test_unit_from_name(cfg):
    from cornerstone.pipeline import unit_from_name
    assert unit_from_name("Unit 4_Storage management_ Disk Schduling.pdf", cfg) == "Unit 4"
    assert unit_from_name("unit05 notes.docx", cfg) == "Unit 5"
    assert unit_from_name("Operating systems.pdf", cfg) is None


def test_ingest_records_unit_and_source_path(cfg, stack, embedder, small_pdf):
    r = run(small_pdf, cfg, stack, embedder)
    d = stack[0].get_document(r.doc_id)
    assert d.unit == "Unit 5" and d.source_path == str(small_pdf.resolve())


def test_duplicate_pages_are_recorded_not_indexed(cfg, stack, embedder):
    store = stack[0]
    r = run(SAMPLES / "Unit 4_Storage management_ Disk Schduling.pdf", cfg, stack, embedder)
    assert r.status == "indexed" and r.pages_duplicate == 23
    assert r.pages_with_text + r.pages_flagged + r.pages_duplicate == r.pages_total == 51
    assert (27, 2) in r.duplicate_pages and all(p != 27 for p, _ in r.flagged_pages)
    chunks = store.get_chunks(store.active_chunk_ids()).values()
    dup = {p for p, _ in r.duplicate_pages}
    assert chunks and all(c.page_start not in dup and c.page_end not in dup for c in chunks)
    s = store.stats()
    assert s["pages_duplicate"] == 23 and s["pages_flagged"] == 0
    assert {"file_name": r.file_name, "page_no": 51, "duplicate_of": 26} in s["duplicate_pages"]


def test_page_images_rendered_and_recorded(cfg, stack, embedder, small_pdf):
    from pathlib import Path
    store = stack[0]
    r = run(small_pdf, cfg, stack, embedder)
    assert r.warnings == []
    rel = store.get_page_image(r.doc_id, 1)
    assert rel == f"{r.doc_id}/p001.png"
    png = Path(cfg["paths"]["image_dir"]) / rel
    assert png.read_bytes()[:8] == b"\x89PNG\r\n\x1a\n" and png.stat().st_size < 1_000_000
    n = store.conn.execute("SELECT COUNT(*) FROM pages WHERE doc_id=? AND image_path IS NOT NULL",
                           (r.doc_id,)).fetchone()[0]
    assert n == r.pages_total == 28                      # flagged pages get images too
    assert store.get_page_image(r.doc_id, 999) is None


def test_page_image_failure_is_a_warning_not_a_crash(cfg, stack, embedder, small_pdf, monkeypatch):
    import pymupdf
    orig = pymupdf.Page.get_pixmap

    def flaky(self, *a, **kw):
        if self.number == 2:
            raise RuntimeError("render failed")
        return orig(self, *a, **kw)
    monkeypatch.setattr(pymupdf.Page, "get_pixmap", flaky)
    r = run(small_pdf, cfg, stack, embedder)
    assert r.status == "indexed" and any("page 3" in w for w in r.warnings)
    assert stack[0].get_page_image(r.doc_id, 3) is None and stack[0].get_page_image(r.doc_id, 4)


def test_failed_ingest_removes_its_images(cfg, stack, embedder, small_pdf, monkeypatch):
    from pathlib import Path
    from cornerstone.index.metadata_store import MetadataStore
    monkeypatch.setattr(MetadataStore, "activate", lambda *a, **k: (_ for _ in ()).throw(OSError("disk full")))
    with pytest.raises(OSError):
        run(small_pdf, cfg, stack, embedder)
    doc_id = stack[0].list_documents("abandoned")[0].doc_id
    assert not (Path(cfg["paths"]["image_dir"]) / str(doc_id)).exists()


def test_images_disabled_renders_nothing(cfg, small_pdf):
    from cornerstone.ingest.page_images import render_page_images
    assert render_page_images(small_pdf, 1, {**cfg, "images": {"enabled": False, "dpi": 110}}) == ({}, [])


# ---- compaction (ARCHITECTURE 10.9)
def _replace_pdf(cfg, stack, embedder, tmp_path):
    """Ingest x.pdf (Unit 5-2), then a different PDF under the same name (Unit 4) -> first version replaced."""
    target = tmp_path / "x.pdf"
    shutil.copy(SAMPLES / "Unit 5_Part 2_System Protection.pdf", target)
    first = run(target, cfg, stack, embedder)
    shutil.copy(SAMPLES / "Unit 4_Storage management_ Disk Schduling.pdf", target)
    second = run(target, cfg, stack, embedder)
    assert second.status == "replaced"
    return first, second


def test_compact_with_nothing_to_do_changes_nothing(cfg, stack, embedder, small_pdf):
    from cornerstone.pipeline import compact
    store, index, ipath, _ = stack
    run(small_pdf, cfg, stack, embedder)
    mtime = ipath.stat().st_mtime_ns
    new, r = compact(cfg, store, index, ipath)
    assert new is index and r["vectors_before"] == r["vectors_after"] and r["chunk_rows_removed"] == 0
    assert ipath.stat().st_mtime_ns == mtime


def test_compact_removes_replaced_vectors_rows_fts_and_images(cfg, stack, embedder, tmp_path):
    from pathlib import Path
    from cornerstone.hybrid import keyword_search
    from cornerstone.pipeline import compact
    store, index, ipath, _ = stack
    first, second = _replace_pdf(cfg, stack, embedder, tmp_path)
    old_imgs = Path(cfg["paths"]["image_dir"]) / str(first.doc_id)
    assert old_imgs.is_dir() and len(index) == first.chunks_created + second.chunks_created
    new, r = compact(cfg, store, index, ipath)
    assert r["vectors_after"] == len(new) == second.chunks_created and r["bytes_after"] < r["bytes_before"]
    assert r["chunk_rows_removed"] == first.chunks_created and r["documents_cleaned"] == 1
    assert store.check_sync(set(new.all_ids().tolist())) == []
    assert keyword_search("setuid", 5, store) == []                       # old version's words are gone
    assert store.conn.execute("SELECT COUNT(*) FROM pages WHERE doc_id=?", (first.doc_id,)).fetchone()[0] == 0
    assert store.get_document(first.doc_id).status == "replaced"           # history kept
    assert not old_imgs.exists()
    reloaded = HNSWIndex.load(ipath, cfg)                                  # what was saved is the compacted index
    assert len(reloaded) == len(new)


def test_compact_refuses_while_an_ingest_is_pending(cfg, stack, embedder, small_pdf):
    from cornerstone import Chunk, Document, IndexOutOfSync, Page
    from cornerstone.pipeline import compact
    store, index, ipath, _ = stack
    run(small_pdf, cfg, stack, embedder)
    store.add_document(Document("p.pdf", "zz", "pdf"), [Page(1, "t")], [Chunk(0, 0, "p.pdf", "t", 0, 1)])
    with pytest.raises(IndexOutOfSync, match="pending"):
        compact(cfg, store, index, ipath)


def test_compact_crash_after_index_swap_is_recoverable(cfg, stack, embedder, tmp_path, monkeypatch):
    from cornerstone.index.metadata_store import MetadataStore
    from cornerstone.pipeline import compact
    store, index, ipath, _ = stack
    _replace_pdf(cfg, stack, embedder, tmp_path)
    monkeypatch.setattr(MetadataStore, "purge_deleted", lambda self: (_ for _ in ()).throw(Crash()))
    with pytest.raises(Crash):
        compact(cfg, store, index, ipath)
    monkeypatch.undo()
    reloaded = HNSWIndex.load(ipath, cfg)                                  # compacted index is on disk
    assert store.check_sync(set(reloaded.all_ids().tolist())) == []        # deleted rows without vectors: fine
    new, r = compact(cfg, store, reloaded, ipath)                          # second run purges the rows
    assert r["chunk_rows_removed"] > 0 and r["vectors_after"] == r["vectors_before"]
